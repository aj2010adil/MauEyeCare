"""
Simple DOCX Generator - Reliable prescription generation
"""
from io import BytesIO
import datetime

def generate_simple_prescription_docx(prescription, doctor_name, patient_name, age, gender, advice, rx_table, recommendations, dosages=None):
    """Generate simple but reliable DOCX prescription"""
    
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, Inches

        # Create document
        doc = Document()

        # Estimate content size to pick a starting font size so document fits one page
        content_preview = ''
        content_preview += f"{patient_name}{age}{gender}"
        if rx_table:
            for eye in ['OD', 'OS']:
                eye_data = rx_table.get(eye, {})
                content_preview += ''.join([str(v) for v in eye_data.values()])
        if prescription:
            for item, qty in prescription.items():
                content_preview += f"{item}{qty}"
                if dosages and item in dosages:
                    dosage_info = dosages[item]
                    content_preview += dosage_info.get('dosage','') + dosage_info.get('timing','')
        if advice:
            content_preview += advice

        length = len(content_preview)
        # base size: 11pt, shrink for larger content
        if length <= 1200:
            base_font_pt = 11
        elif length <= 2000:
            base_font_pt = 10
        elif length <= 3000:
            base_font_pt = 9
        else:
            base_font_pt = 8
        
        # Reduce margins to maximize printable area and set default fonts
        section = doc.sections[0]
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

        # Apply default normal style
        try:
            normal_style = doc.styles['Normal']
            normal_style.font.name = 'Arial'
            normal_style.font.size = Pt(base_font_pt)
        except Exception:
            pass

        # Header
        header = doc.add_heading('MauEyeCare Optical Center', 0)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.runs[0].font.size = Pt(base_font_pt + 3)
        
        # Clinic info
        clinic_para = doc.add_paragraph()
        clinic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = clinic_para.add_run('Dr. Danish - Eye Care Specialist\n')
        r.bold = True
        r.font.size = Pt(base_font_pt)
        r2 = clinic_para.add_run('Phone: +91 92356-47410 | Email: info@maueyecare.com')
        r2.font.size = Pt(base_font_pt)
        
        # Date
        timestamp = datetime.datetime.now()
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        d1 = date_para.add_run(f'Date: {timestamp.strftime("%d/%m/%Y")}\n')
        d1.font.size = Pt(base_font_pt)
        d2 = date_para.add_run(f'Prescription No: RX-{timestamp.strftime("%Y%m%d%H%M")}')
        d2.font.size = Pt(base_font_pt)
        
        # Patient info
        p1 = doc.add_paragraph(f'Patient Name: {patient_name}')
        p1.runs[0].font.size = Pt(base_font_pt)
        p2 = doc.add_paragraph(f'Age: {age} years | Gender: {gender}')
        p2.runs[0].font.size = Pt(base_font_pt)
        
        # RX Details
        if rx_table:
            hdr = doc.add_heading('Prescription Details:', level=2)
            try:
                hdr.runs[0].font.size = Pt(base_font_pt + 1)
            except Exception:
                pass
            for eye in ['OD', 'OS']:
                eye_data = rx_table.get(eye, {})
                if eye_data.get('Sphere'):
                    rx_line = f"{eye}: "
                    if eye_data.get('Sphere'):
                        rx_line += f"SPH {eye_data['Sphere']} "
                    if eye_data.get('Cylinder'):
                        rx_line += f"CYL {eye_data['Cylinder']} "
                    if eye_data.get('Axis'):
                        rx_line += f"AXIS {eye_data['Axis']}"
                    p = doc.add_paragraph(rx_line)
                    try:
                        p.runs[0].font.size = Pt(base_font_pt)
                    except Exception:
                        pass
        
        # Medicines
        if prescription:
            hdr = doc.add_heading('Prescribed Medications:', level=2)
            try:
                hdr.runs[0].font.size = Pt(base_font_pt + 1)
            except Exception:
                pass
            for item, qty in prescription.items():
                med_para = doc.add_paragraph()
                r = med_para.add_run(f'• {item} - Quantity: {qty}')
                r.bold = True
                r.font.size = Pt(base_font_pt)

                if dosages and item in dosages:
                    dosage_info = dosages[item]
                    d1 = med_para.add_run(f'\n  Dosage: {dosage_info.get("dosage", "As directed")}')
                    d1.font.size = Pt(base_font_pt)
                    d2 = med_para.add_run(f'\n  Timing: {dosage_info.get("timing", "As directed")}')
                    d2.font.size = Pt(base_font_pt)
        
        # Advice
        if advice:
            hdr = doc.add_heading('Doctor\'s Advice:', level=2)
            try:
                hdr.runs[0].font.size = Pt(base_font_pt + 1)
            except Exception:
                pass
            a_para = doc.add_paragraph(advice)
            try:
                a_para.runs[0].font.size = Pt(base_font_pt)
            except Exception:
                pass
        
        # Instructions
        doc.add_heading('Instructions:', level=2)
        instructions = [
            'Take medications as prescribed',
            'Follow up if symptoms persist',
            'Avoid rubbing eyes',
            'Contact clinic for emergencies'
        ]
        for instruction in instructions:
            doc.add_paragraph(f'• {instruction}')
        
        # Signature
        signature_para = doc.add_paragraph()
        signature_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sig = signature_para.add_run('\n\nDr. Danish\nEye Care Specialist')
        try:
            sig.font.size = Pt(base_font_pt)
        except Exception:
            pass
        
        # Save to buffer
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
        
    except Exception as e:
        # Fallback to text format
        return generate_text_prescription(prescription, doctor_name, patient_name, age, gender, advice, rx_table, recommendations, dosages)

def generate_text_prescription(prescription, doctor_name, patient_name, age, gender, advice, rx_table, recommendations, dosages=None):
    """Generate text-based prescription as fallback"""
    
    timestamp = datetime.datetime.now()
    
    content = f"""MauEyeCare Optical Center
Dr. Danish - Eye Care Specialist
Phone: +91 92356-47410 | Email: info@maueyecare.com

Date: {timestamp.strftime("%d/%m/%Y")}
Prescription No: RX-{timestamp.strftime("%Y%m%d%H%M")}

Patient Name: {patient_name}
Age: {age} years | Gender: {gender}

"""
    
    # RX Details
    if rx_table:
        content += "Prescription Details:\n"
        for eye in ['OD', 'OS']:
            eye_data = rx_table.get(eye, {})
            if eye_data.get('Sphere'):
                rx_line = f"{eye}: "
                if eye_data.get('Sphere'):
                    rx_line += f"SPH {eye_data['Sphere']} "
                if eye_data.get('Cylinder'):
                    rx_line += f"CYL {eye_data['Cylinder']} "
                if eye_data.get('Axis'):
                    rx_line += f"AXIS {eye_data['Axis']}"
                content += rx_line + "\n"
        content += "\n"
    
    # Medicines
    if prescription:
        content += "Prescribed Medications:\n"
        for item, qty in prescription.items():
            content += f"• {item} - Quantity: {qty}\n"
            if dosages and item in dosages:
                dosage_info = dosages[item]
                content += f"  Dosage: {dosage_info.get('dosage', 'As directed')}\n"
                content += f"  Timing: {dosage_info.get('timing', 'As directed')}\n"
        content += "\n"
    
    # Advice
    if advice:
        content += f"Doctor's Advice:\n{advice}\n\n"
    
    # Instructions
    content += """Instructions:
• Take medications as prescribed
• Follow up if symptoms persist
• Avoid rubbing eyes
• Contact clinic for emergencies

                                Dr. Danish
                                Eye Care Specialist
"""
    
    return content.encode('utf-8')